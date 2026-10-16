"""The one-page resume format: header with the target title, bulleted career
summary, approved alternative job titles, relevance-selected roles, a skills
section only when the posting calls for one, bold key phrases, clean document
metadata, and length fitting to one page."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from resume_gen.guard import clean_headline, enforce
from resume_gen.models import Contact, ExperienceItem, Resume
from resume_gen.render.docx_renderer import _date_range, render_cover_letter, render_resume
from resume_gen.render.length import fit_length
from resume_gen.models import CoverLetter

PROFILE = {
    "full_name": "Keerthivasan Kalaivanan",
    "resume_name": "Keerthivasan",
    "preferred_name": "Keerthi",
    "contact": {
        "email": "k@example.com", "phone": "555-0100", "location": "London, ON",
        "links": [{"label": "LinkedIn", "url": "https://linkedin.com/in/kkvasan/"}],
    },
    "skills": {"data_analytics": ["SQL", "Excel", "Power BI"], "tools_other": ["Git"]},
    "experience": [
        {"company": "K2 Digital Media", "role": "Freelance Full-Stack Developer and Owner",
         "titles": ["Founder and Project Lead"], "start": "2023", "end": "Present",
         "always_include": True,
         "facts": ["Lead project kick-offs with clients and vendors and set project timelines.",
                   "Build Excel and Power BI reports for client decision makers.",
                   "Manage the critical path across several client projects at once."]},
        {"company": "Certify Technologies", "role": "Associate Developer (Application and IT Support)",
         "titles": ["Associate Developer"], "start": "Nov 2021", "end": "Jul 2023",
         "always_include": True,
         "facts": ["Maintained a 95% first-contact resolution rate on OS and account issues.",
                   "Partnered with QA and product teams on release sign-off."]},
        {"company": "Best Buy", "role": "Sales Associate, Computing", "start": "Sep 2023",
         "end": "Dec 2024", "dates": "Sep to Dec 2023, Aug to Dec 2024",
         "facts": ["Merchandised product on a high-volume retail floor."]},
        {"company": "Haply Robotics (Fanshawe placement)", "role": "Research Assistant (Real-Time Systems)",
         "start": "Aug 2024", "end": "Dec 2024",
         "facts": ["Built C# and C++ interop between Unity and CHAI 3D."]},
    ],
    "education": [{"institution": "Fanshawe College, London, ON",
                   "credential": "Game Development: Advanced Programming", "year": "2023 to 2024"}],
    "certifications": [],
}


def _resume(**kw) -> Resume:
    base = dict(
        fullName="Keerthivasan", headline="", contact=Contact(email="", phone="", location=""),
        summaryPoints=["Runs **client projects from kick-off to launch** across timelines and vendors.",
                       "Builds Excel and Power BI reports for decision makers.",
                       "Brings a B.Com and retail merchandising experience."],
        skills=[],
        experience=[
            ExperienceItem(company="K2 Digital Media", role="Founder and Project Lead", start="2023",
                           end="Present", bullets=["Lead **project kick-offs** with clients and vendors."]),
            ExperienceItem(company="Best Buy", role="Sales Associate, Computing", start="", end="",
                           bullets=["Merchandised product on a high-volume retail floor."]),
        ],
    )
    base.update(kw)
    return Resume(**base)


class GuardFormatTests(unittest.TestCase):
    def test_headline_is_the_posting_title_not_the_model_text(self):
        out, _ = enforce(_resume(headline="Senior Wizard"), PROFILE,
                         target_title="Associate Product Manager, Value Creation (Contract) - Req 12345")
        self.assertEqual(out.headline, "Associate Product Manager, Value Creation")

    def test_clean_headline_drops_dashes_and_brackets(self):
        self.assertEqual(clean_headline("QA Analyst – Remote (12 month contract)"), "QA Analyst, Remote")

    def test_approved_alternative_title_is_kept_and_others_are_reset(self):
        r = _resume(experience=[
            ExperienceItem(company="K2 Digital Media", role="founder and project lead", start="", end="",
                           bullets=["Lead project kick-offs with clients."]),
            ExperienceItem(company="Certify Technologies", role="Senior Software Engineer", start="", end="",
                           bullets=["Partnered with QA on release sign-off."]),
        ])
        out, qa = enforce(r, PROFILE, target_title="X")
        roles = {e.company: e.role for e in out.experience}
        self.assertEqual(roles["K2 Digital Media"], "Founder and Project Lead")
        self.assertEqual(roles["Certify Technologies"], "Associate Developer (Application and IT Support)")
        self.assertTrue(any("Senior Software Engineer" in c for f in qa["experience_fixed"] for c in f["changes"]))

    def test_always_include_roles_are_added_and_other_omissions_are_reported(self):
        out, qa = enforce(_resume(), PROFILE, target_title="X")
        companies = [e.company for e in out.experience]
        self.assertIn("Certify Technologies", companies)            # always_include, re-added
        self.assertNotIn("Haply Robotics (Fanshawe placement)", companies)
        self.assertEqual(qa["roles_omitted"], ["Haply Robotics (Fanshawe placement)"])
        self.assertEqual(qa["experience_added_from_profile"], ["Certify Technologies"])

    def test_dates_override_comes_from_the_profile(self):
        out, _ = enforce(_resume(), PROFILE, target_title="X")
        best_buy = next(e for e in out.experience if e.company == "Best Buy")
        self.assertEqual(best_buy.dates, "Sep to Dec 2023, Aug to Dec 2024")

    def test_empty_skills_stay_empty_but_thin_skills_are_backfilled(self):
        out, _ = enforce(_resume(skills=[]), PROFILE, target_title="X")
        self.assertEqual(out.skills, [])
        out, qa = enforce(_resume(skills=["SQL"]), PROFILE, target_title="X")
        self.assertIn("SQL", out.skills)
        self.assertGreater(len(out.skills), 1)

    def test_paragraph_summary_becomes_three_points(self):
        r = _resume(summaryPoints=[], summary="One thing. Two things. Three things. Four things.")
        out, qa = enforce(r, PROFILE, target_title="X")
        self.assertEqual(out.summaryPoints, ["One thing.", "Two things.", "Three things."])
        self.assertEqual(out.summary, "")
        self.assertTrue(qa["summary_split_into_points"])

    def test_location_only_on_local_roles(self):
        out, _ = enforce(_resume(), PROFILE, target_title="X", target_location="London, ON")
        self.assertEqual(out.contact.location, "London, ON")
        out, _ = enforce(_resume(), PROFILE, target_title="X", target_location="Toronto, ON")
        self.assertEqual(out.contact.location, "")


class RenderFormatTests(unittest.TestCase):
    def _render(self, resume: Resume):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.docx"
            render_resume(resume, path, PROFILE)
            return Document(path)

    def test_header_name_preferred_name_and_title(self):
        out, _ = enforce(_resume(), PROFILE, target_title="Associate Product Manager",
                         target_location="Brampton, ON")
        doc = self._render(out)
        self.assertEqual(doc.paragraphs[0].text, "KEERTHIVASAN  (Keerthi)\tAssociate Product Manager")
        self.assertEqual(doc.paragraphs[1].text, "k@example.com  |  555-0100  |  linkedin.com/in/kkvasan")

    def test_key_phrases_render_bold_without_markers(self):
        doc = self._render(_resume())
        bullet = next(p for p in doc.paragraphs if "kick-offs" in p.text)
        self.assertNotIn("**", bullet.text)
        self.assertEqual([r.text for r in bullet.runs if r.bold], ["project kick-offs"])

    def test_no_skills_heading_when_skills_are_empty(self):
        headings = [p.text for p in self._render(_resume()).paragraphs if p.style.name == "Heading 2"]
        self.assertEqual(headings, ["CAREER SUMMARY", "PROFESSIONAL EXPERIENCE"])

    def test_metadata_carries_no_tool_trace(self):
        doc = self._render(_resume())
        cp = doc.core_properties
        self.assertEqual(cp.author, "Keerthivasan Kalaivanan")
        self.assertEqual(cp.comments, "")
        self.assertNotIn("python-docx", f"{cp.author}{cp.comments}{cp.last_modified_by}")
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "cl.docx"
            render_cover_letter(CoverLetter(fullName="Keerthivasan Kalaivanan", contactLine="x",
                                            greeting="Hi,", body=["One."], signature="Keerthi"), path)
            cp = Document(path).core_properties
            self.assertNotIn("python-docx", f"{cp.author}{cp.comments}")

    def test_date_ranges_use_to_and_compress_one_year(self):
        self.assertEqual(_date_range("Aug 2024", "Nov 2024"), "Aug to Nov 2024")
        self.assertEqual(_date_range("2023", "present"), "2023 to Present")
        self.assertEqual(_date_range("Nov 2021", "Jul 2023"), "Nov 2021 to Jul 2023")


class LengthFitTests(unittest.TestCase):
    def test_short_resume_fits_one_page_untrimmed(self):
        r, k, rep = fit_length(_resume(), PROFILE)
        self.assertEqual(rep["pages_target"], 1)
        self.assertEqual(rep["trimmed"], [])
        self.assertEqual(k, 1.30)                     # loosest density: fill the page

    LONG = ("Built and maintained weekly client reporting in Excel and Power BI for business "
            "owners across several accounts, consolidating campaign results into one view. ") * 2

    def test_overlong_resume_is_trimmed_to_one_page_before_going_to_two(self):
        roles = [ExperienceItem(company=f"Co {i}", role="Analyst", start="2020", end="2021",
                                bullets=[self.LONG] * 5) for i in range(5)]
        r, k, rep = fit_length(_resume(experience=roles), PROFILE)
        self.assertEqual(rep["pages_target"], 1)
        self.assertTrue(rep["trimmed"])
        self.assertLessEqual(rep["estimated_fill"], 1.0)
        self.assertGreaterEqual(len(r.experience), 3)
        self.assertGreaterEqual(len(r.experience[0].bullets), 3)

    def test_two_pages_only_when_trimming_cannot_fit_one(self):
        # Every role is always-included, so trimming stops at 3 lead bullets and
        # 2 per other role, which still overflows one page.
        roles = [ExperienceItem(company=("K2 Digital Media" if i % 2 else "Certify Technologies"),
                                role="Analyst", start="2020", end="2021", bullets=[self.LONG * 2] * 5)
                 for i in range(6)]
        r, k, rep = fit_length(_resume(experience=roles), PROFILE)
        self.assertEqual(rep["pages_target"], 2)
        self.assertEqual(sum(len(e.bullets) for e in r.experience), 30)   # untrimmed
        self.assertIn("page_two_fill_ok", rep)

if __name__ == "__main__":
    unittest.main()
